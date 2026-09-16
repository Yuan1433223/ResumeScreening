"""批量处理与结果导出"""
import json
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

from .workflow import run_single
from .config import RESUME_DIR, OUTPUT_DIR, JOB_TITLE


def list_resumes(resume_dir: Path = RESUME_DIR) -> List[Path]:
    """列出所有简历文件"""
    exts = {".pdf", ".jpg", ".jpeg", ".png", ".bmp"}
    files = []
    for p in resume_dir.iterdir():
        if p.is_file() and p.suffix.lower() in exts:
            files.append(p)
    return sorted(files, key=lambda x: x.name)


def process_batch(limit: int = None, resume_dir: Path = RESUME_DIR) -> Dict[str, Any]:
    """批量处理简历"""
    files = list_resumes(resume_dir)
    if limit:
        files = files[:limit]

    reports = []
    for idx, fp in enumerate(files, 1):
        print(f"[{idx}/{len(files)}] 处理: {fp.name}")
        try:
            # 不传 thread_id：run_single 以稳定 document_id 作为 thread（跨批次稳定，
            # 不再使用 resume_{idx} 位置编号，避免排序变化导致线程串用）
            report = run_single(str(fp), JOB_TITLE)
        except Exception as e:
            report = {
                "岗位": JOB_TITLE,
                "候选人": fp.stem,
                "文件名": fp.name,
                "解析状态": "failed",
                "解析错误": str(e),
                "机器建议": "needs_review",
            }
        reports.append(report)

    passed = [r for r in reports if r.get("机器建议") == "recommend"]
    scheduled = [r for r in reports if r.get("面试安排")]
    return {
        "生成时间": datetime.now().isoformat(),
        "岗位": JOB_TITLE,
        "处理总数": len(reports),
        "推荐数": len(passed),
        "安排面试数": len(scheduled),
        "报告列表": reports,
    }


def save_results(batch_result: Dict[str, Any]) -> Dict[str, str]:
    """保存结果到 JSON/Excel(openpyxl)/Markdown（不依赖 pandas）"""
    OUTPUT_DIR.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    json_path = OUTPUT_DIR / f"result_{timestamp}.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(batch_result, f, ensure_ascii=False, indent=2)

    # Excel 导出（openpyxl，已声明的依赖）
    excel_path = None
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "简历筛选"
        headers = [
            "候选人", "文件名", "解析状态", "硬条件结果", "综合评分",
            "评分完整", "机器建议", "面试安排", "复核意见",
        ]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")
        for r in batch_result["报告列表"]:
            ws.append([
                r.get("候选人", ""),
                r.get("文件名", ""),
                r.get("解析状态", ""),
                "; ".join(f"{k}={v}" for k, v in (r.get("硬条件结果") or {}).items()),
                r.get("综合评分"),
                r.get("评分完整", ""),
                r.get("机器建议", ""),
                r.get("面试安排", ""),
                "; ".join(r.get("复核意见", [])),
            ])
        # 列宽
        for col, w in zip("ABCDEFGHI", [14, 30, 10, 40, 10, 8, 14, 10, 40]):
            ws.column_dimensions[col].width = w
        excel_path = OUTPUT_DIR / f"result_{timestamp}.xlsx"
        wb.save(str(excel_path))
    except Exception as e:
        excel_path = None
        print(f"Excel 导出失败: {e}")

    # Markdown 报告
    md_path = OUTPUT_DIR / f"result_{timestamp}.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# 简历筛选报告\n\n")
        f.write(f"- 岗位：{batch_result['岗位']}\n")
        f.write(f"- 处理总数：{batch_result['处理总数']}\n")
        f.write(f"- 推荐数：{batch_result['推荐数']}\n")
        f.write(f"- 安排面试数：{batch_result['安排面试数']}\n")
        f.write(f"- 生成时间：{batch_result['生成时间']}\n\n")
        f.write("## 候选人清单\n\n")
        f.write("| 候选人 | 解析状态 | 硬条件 | 评分 | 建议 | 面试 | 文件名 |\n")
        f.write("| --- | --- | --- | --- | --- | --- | --- |\n")
        for r in batch_result["报告列表"]:
            hr = r.get("硬条件结果") or {}
            hr_txt = "; ".join(f"{k}:{v}" for k, v in hr.items()) if hr else "-"
            f.write(
                f"| {r.get('候选人')} | {r.get('解析状态')} | {hr_txt} | "
                f"{r.get('综合评分')} | {r.get('机器建议')} | {r.get('面试安排')} | {r.get('文件名')} |\n"
            )

    return {
        "json": str(json_path),
        "excel": str(excel_path) if excel_path else "",
        "markdown": str(md_path),
    }
